import json

from fastapi import UploadFile
from sqlalchemy.orm import Session

from app.generators.certificate_generator import generate_certificate_pdf_for_job
from app.repositories.job_repository import JobRepository
from app.storage import storage
from app.repositories.template_repository import TemplateRepository
from app.validators.excel_validator import validate_excel_for_template


class GenerationService:
    def __init__(self, db: Session):
        self.db = db
        self.job_repository = JobRepository(db)
        self.template_repository = TemplateRepository(db)

    def create_job(self, template_id: int):
        template = self.template_repository.get_by_id(template_id)
        if template is None:
            raise ValueError("Template not found")
        job = self.job_repository.create_job(template_id)
        return job

    def validate_upload(self, template_id: int, file: UploadFile):
        template = self.template_repository.get_by_id(template_id)
        if template is None:
            raise ValueError("Template not found")
        contents = file.file.read()
        file.file.seek(0)
        import io
        validation = validate_excel_for_template(io.BytesIO(contents), template)
        return validation

    def process_generation(self, job_id: int, file: UploadFile):
        job = self.job_repository.get_job_by_id(job_id)
        if job is None:
            raise ValueError("Generation job not found")
        template = self.template_repository.get_by_id(job.template_id)
        if template is None:
            raise ValueError("Template not found")

        import io
        file_bytes = file.file.read()
        file.file.seek(0)
        validation = validate_excel_for_template(io.BytesIO(file_bytes), template)
        if not validation["valid"]:
            self.job_repository.update_job_status(job, "FAILED", "Validation failed before generation")
            return job

        data_rows = self._read_rows(file_bytes)
        job.total_count = len(data_rows)
        job.pending_count = len(data_rows)
        job.status = "PROCESSING"
        self.db.commit()

        for row_number, row_data in data_rows:
            certificate = self.job_repository.add_certificate(job.id, row_number, row_data)
            row_data = {key: value for key, value in row_data.items() if value is not None}
            try:
                job.processing_count += 1
                job.pending_count = max(job.pending_count - 1, 0)
                self.db.commit()
                pdf_bytes = generate_certificate_pdf_for_job(template, row_data, job.id, row_number)
                certificate_name = f"certificate_{row_number}.pdf"
                file_key = storage.upload_bytes(
                    f"certificates/{job.id}/{certificate_name}",
                    pdf_bytes,
                    "application/pdf",
                )
                self.job_repository.update_certificate(certificate, "SUCCESS", file_path=file_key, certificate_name=certificate_name)
                job.success_count += 1
            except Exception as exc:  # pragma: no cover - runtime safeguard
                self.job_repository.update_certificate(certificate, "FAILED", error_message=str(exc))
                job.failed_count += 1
            finally:
                job.processing_count = max(job.processing_count - 1, 0)
                self.db.commit()

        if job.failed_count > 0:
            job.status = "COMPLETED_WITH_ERRORS"
        else:
            job.status = "COMPLETED"
        job.completed_at = job.completed_at or __import__("datetime").datetime.utcnow()
        self.db.commit()
        return job

    def _read_rows(self, file_bytes):
        import io
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = workbook.active
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return []

        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        cleaned_headers = [header.lower().replace(" ", "_") for header in headers]
        items = []
        for row_index, row in enumerate(rows[1:], start=1):
            row_data = {}
            for header, value in zip(cleaned_headers, row):
                row_data[header] = "" if value is None else str(value).strip()
            items.append((row_index, row_data))
        return items

    def get_job_details(self, job_id: int):
        job = self.job_repository.get_job_by_id(job_id)
        if job is None:
            raise ValueError("Job not found")
        return job

    def get_certificates(self, job_id: int):
        job = self.job_repository.get_job_by_id(job_id)
        if job is None:
            raise ValueError("Job not found")
        return self.job_repository.get_certificates_by_job(job_id)

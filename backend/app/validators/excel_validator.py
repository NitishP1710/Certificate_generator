import re
from typing import Any

from openpyxl import load_workbook

PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")


def extract_placeholders(content: str):
    return list(dict.fromkeys(match.group(1) for match in PLACEHOLDER_PATTERN.finditer(content)))


def normalize_column_name(value: Any) -> str:
    return str(value).strip().lower().replace(" ", "_")


def _clean_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def validate_excel_for_template(file_obj, template) -> dict:
    errors = []
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    ws = workbook.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {
            "valid": False,
            "total_rows": 0,
            "valid_rows": 0,
            "invalid_rows": 0,
            "errors": [{"row": None, "field": None, "message": "Excel file is empty."}],
            "missing_columns": [],
            "unexpected_columns": [],
            "file_name": getattr(file_obj, "filename", "uploaded.xlsx"),
        }

    headers = [str(item).strip() if item is not None else "" for item in rows[0]]
    normalized_headers = [normalize_column_name(item) for item in headers]

    template_placeholders = [p.placeholder_name.lower() for p in template.placeholders]
    required_placeholders = [p.placeholder_name.lower() for p in template.placeholders if p.required]

    missing_columns = [name for name in required_placeholders if name not in normalized_headers]
    unexpected_columns = [name for name in normalized_headers if name not in template_placeholders]

    for column_name in missing_columns:
        errors.append({"row": None, "field": column_name, "message": f"Missing required column: {column_name}"})
    for column_name in unexpected_columns:
        errors.append({"row": None, "field": column_name, "message": f"Unexpected column: {column_name}"})

    total_data_rows = max(len(rows) - 1, 0)
    valid_rows = 0
    invalid_rows = 0

    for row_index, row in enumerate(rows[1:], start=2):
        record = {}
        for header, raw_value in zip(headers, row):
            if header is None:
                continue
            key = normalize_column_name(header)
            record[key] = _clean_value(raw_value)

        row_has_error = False
        for placeholder in template.placeholders:
            field_name = placeholder.placeholder_name.lower()
            value = record.get(field_name, "")
            if placeholder.required and not value:
                errors.append({
                    "row": row_index - 1,
                    "field": placeholder.placeholder_name,
                    "message": "Required value is missing",
                })
                row_has_error = True

        if row_has_error:
            invalid_rows += 1
        else:
            valid_rows += 1

    final_valid = not errors and total_data_rows > 0
    return {
        "valid": final_valid,
        "total_rows": total_data_rows,
        "valid_rows": valid_rows,
        "invalid_rows": invalid_rows,
        "errors": errors,
        "missing_columns": missing_columns,
        "unexpected_columns": unexpected_columns,
        "file_name": getattr(file_obj, "filename", "uploaded.xlsx"),
    }
